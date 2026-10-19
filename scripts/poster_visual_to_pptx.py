"""Editable PowerPoint copy of the VISUAL-FIRST review poster, converted from its PDF (2026-10-19).

  python scripts/poster_visual_to_pptx.py
      -> docs/summit_20261023_reframe/poster_visual_final_review.pptx

One A0 slide:
- **Figures** are the original PNGs from docs/summit_20261023_reframe/figures/, matched by pixel size and placed where
  the PDF places them.
- **TikZ diagrams** are clustered stroke by stroke, so each diagram is its own region. Each region is rendered at
  300 dpi and inserted as a picture; its labels are inside the picture.
- **Every other text line** is an editable text box at its position, size, weight and colour (Arial for Helvetica).
- **The baseline box outline** becomes an editable rectangle.

It is adapted from scripts/poster_to_pptx.py, which converts the submitted poster and is left unchanged.
"""
import io
import os

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIR = os.path.join(HERE, "docs", "summit_20261023_reframe")
EMU = 12700                                           # per PDF point
DIAGRAM_W = (2.4, 3.6)                                 # box borders 2.59, arrows 2.79, the adequacy box 3.39
RULE_W = 5.0                                           # title and block rules: filled bars
OUTLINE_W = (1.4, 1.8)                                 # the baseline box


def _clusters(rects, pad=22):
    """Merge rectangles whose padded copies touch, until stable."""
    out = [fitz.Rect(r) for r in rects]
    merged = True
    while merged:
        merged = False
        res = []
        for r in out:
            for i, q in enumerate(res):
                if (q + (-pad, -pad, pad, pad)).intersects(r):
                    res[i] = q | r
                    merged = True
                    break
            else:
                res.append(fitz.Rect(r))
        out = res
    return out


def main():
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--pdf", default=os.path.join(DIR, "poster_visual_final_review.pdf"))
    ap.add_argument("--out", default=os.path.join(DIR, "poster_visual_final_review.pptx"))
    a = ap.parse_args()
    page = fitz.open(a.pdf)[0]
    W, H = page.rect.width, page.rect.height
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(W * EMU)), Emu(int(H * EMU))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0xFC, 0xFC, 0xFB)
    draws = page.get_drawings()

    regions = [r + (-8, -8, 8, 8) for r in _clusters(
        [d["rect"] for d in draws if DIAGRAM_W[0] <= (d.get("width") or 0) <= DIAGRAM_W[1]])]
    for r in regions:
        pix = page.get_pixmap(clip=r, dpi=300)
        slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), Emu(int(r.x0 * EMU)), Emu(int(r.y0 * EMU)),
                                 Emu(int(r.width * EMU)), Emu(int(r.height * EMU)))
    for d in draws:
        w = d.get("width") or 0
        r = d["rect"]
        if w > RULE_W and d.get("color"):
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(r.x0 * EMU)), Emu(int((r.y0 - w / 2) * EMU)),
                                         Emu(int(max(r.width, w) * EMU)), Emu(int((r.height + w) * EMU)))
            bar.fill.solid()
            bar.fill.fore_color.rgb = RGBColor(*(int(round(v * 255)) for v in d["color"]))
            bar.line.fill.background()
        elif OUTLINE_W[0] <= w <= OUTLINE_W[1] and d.get("color"):
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(r.x0 * EMU)), Emu(int(r.y0 * EMU)),
                                         Emu(int(r.width * EMU)), Emu(int(r.height * EMU)))
            box.fill.background()
            box.line.color.rgb = RGBColor(*(int(round(v * 255)) for v in d["color"]))
            box.line.width = Pt(w)
            box.adjustments[0] = 0.02

    figs = {}
    for name in os.listdir(os.path.join(DIR, "figures")):
        if name.endswith(".png"):
            pm = fitz.Pixmap(os.path.join(DIR, "figures", name))
            figs[(pm.width, pm.height)] = os.path.join(DIR, "figures", name)
    placed = []
    for img in page.get_images(full=True):
        xref, w, h = img[0], img[2], img[3]
        for rect in page.get_image_rects(xref):
            path = figs[(w, h)]
            slide.shapes.add_picture(path, Emu(int(rect.x0 * EMU)), Emu(int(rect.y0 * EMU)),
                                     Emu(int(rect.width * EMU)), Emu(int(rect.height * EMU)))
            placed.append(os.path.basename(path))

    lines = 0
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            box = fitz.Rect(line["bbox"])
            if any(r.contains(box) or (r.intersects(box) and (r & box).get_area() > 0.5 * box.get_area())
                   for r in regions):
                continue
            spans = [s for s in line["spans"] if s["text"].strip()]
            if not spans:
                continue
            tb = slide.shapes.add_textbox(Emu(int(box.x0 * EMU)), Emu(int(box.y0 * EMU)),
                                          Emu(int((box.width + 4) * EMU)), Emu(int(box.height * EMU)))
            tf = tb.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.word_wrap = False
            para = tf.paragraphs[0]
            for s in line["spans"]:
                run = para.add_run()
                run.text = s["text"]
                run.font.size = Pt(round(s["size"], 1))
                run.font.name = "Courier New" if "Courier" in s["font"] or "cmtt" in s["font"].lower() else "Arial"
                run.font.bold = bool(s["flags"] & 16) or "Bold" in s["font"]
                run.font.italic = bool(s["flags"] & 2) or "Oblique" in s["font"] or "Italic" in s["font"]
                c = s["color"]
                run.font.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
            lines += 1
    prs.core_properties.title = "Learning What to Change (CORA, ARC Prize Research Summit 2026), visual review draft"
    prs.save(a.out)
    print({"slide_cm": [round(W / 72 * 2.54, 1), round(H / 72 * 2.54, 1)], "diagram_regions": len(regions),
           "figures": placed, "text_lines": lines, "out": a.out})


if __name__ == "__main__":
    main()
