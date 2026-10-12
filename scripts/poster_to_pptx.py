"""Editable PowerPoint copy of the summit poster, converted from poster.pdf.

  python scripts/poster_to_pptx.py   -> docs/summit_poster/poster.pptx

One A0 slide. Every text line of the PDF becomes an editable text box at its
position, size, weight and colour (Arial for Helvetica). Each figure is the
original PNG from docs/summit_poster/figures/ (matched by pixel size), placed
where the PDF places it. The two TikZ diagrams are rendered from the PDF at
300 dpi and inserted as pictures; their text is inside the picture.
"""
import io
import os

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
POSTER = os.path.join(HERE, "docs", "summit_poster")
EMU = 12700                                           # per PDF point


def main():
    page = fitz.open(os.path.join(POSTER, "poster.pdf"))[0]
    W, H = page.rect.width, page.rect.height
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(W * EMU)), Emu(int(H * EMU))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0xFC, 0xFC, 0xFB)

    # diagrams: clusters of thick vector strokes (node borders, arrows), one per column
    cols = [(0, W / 3), (W / 3, 2 * W / 3), (2 * W / 3, W)]
    regions = []
    for x0, x1 in cols:
        rs = [d["rect"] for d in page.get_drawings() if 1.5 <= (d.get("width") or 0) <= 3.0
              and x0 <= (d["rect"].x0 + d["rect"].x1) / 2 < x1]
        if rs:
            r = fitz.Rect(rs[0])
            for q in rs[1:]:
                r |= q
            regions.append(r + (-6, -6, 6, 6))
    for r in regions:
        pix = page.get_pixmap(clip=r, dpi=300)
        slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), Emu(int(r.x0 * EMU)), Emu(int(r.y0 * EMU)),
                                 Emu(int(r.width * EMU)), Emu(int(r.height * EMU)))

    # block-title and title rules (thick strokes outside the diagrams), as filled bars
    for d in page.get_drawings():
        w = d.get("width") or 0
        if w > 3.0 and d.get("color"):
            r = d["rect"]
            bar = slide.shapes.add_shape(1, Emu(int(r.x0 * EMU)), Emu(int((r.y0 - w / 2) * EMU)),
                                         Emu(int(max(r.width, w) * EMU)), Emu(int((r.height + w) * EMU)))
            bar.fill.solid()
            bar.fill.fore_color.rgb = RGBColor(*(int(round(v * 255)) for v in d["color"]))
            bar.line.fill.background()

    # figures: the original PNGs, matched by pixel size
    figs = {}
    for name in os.listdir(os.path.join(POSTER, "figures")):
        if name.endswith(".png"):
            pm = fitz.Pixmap(os.path.join(POSTER, "figures", name))
            figs[(pm.width, pm.height)] = os.path.join(POSTER, "figures", name)
    placed = 0
    for img in page.get_images(full=True):
        xref, w, h = img[0], img[2], img[3]
        for rect in page.get_image_rects(xref):
            path = figs[(w, h)]
            slide.shapes.add_picture(path, Emu(int(rect.x0 * EMU)), Emu(int(rect.y0 * EMU)),
                                     Emu(int(rect.width * EMU)), Emu(int(rect.height * EMU)))
            placed += 1

    # text: one editable box per line, outside the diagram pictures
    lines = 0
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            box = fitz.Rect(line["bbox"])
            if any(r.contains(box) or r.intersects(box) and (r & box).get_area() > 0.5 * box.get_area()
                   for r in regions):
                continue
            spans = [s for s in line["spans"] if s["text"].strip()]
            if not spans:
                continue
            tb = slide.shapes.add_textbox(Emu(int(box.x0 * EMU)), Emu(int(box.y0 * EMU)),
                                          Emu(int((box.width + 4) * EMU)), Emu(int(box.height * EMU)))
            tf = tb.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.word_wrap = False
            para = tf.paragraphs[0]
            for s in line["spans"]:
                run = para.add_run()
                run.text = s["text"]
                run.font.size = Pt(round(s["size"], 1))
                run.font.name = "Courier New" if "Courier" in s["font"] or "cmtt" in s["font"].lower() else "Arial"
                run.font.bold = bool(s["flags"] & 16) or "Bold" in s["font"]
                run.font.italic = bool(s["flags"] & 2) or "Oblique" in s["font"] or "Italic" in s["font"]
                c = s["color"]
                run.font.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
            lines += 1
    prs.core_properties.title = "When Search Is Not Enough (CORA, ARC Prize Research Summit 2026)"
    out = os.path.join(POSTER, "poster.pptx")
    prs.save(out)
    print({"slide_cm": [round(W / 72 * 2.54, 1), round(H / 72 * 2.54, 1)], "diagrams": len(regions),
           "figures": placed, "text_lines": lines, "out": out})


if __name__ == "__main__":
    main()
